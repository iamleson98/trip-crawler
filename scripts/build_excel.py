#!/usr/bin/env python3
"""Build a multi-sheet Excel workbook from the N×N crawl results.

Reads /tmp/nxn_out/manifest.json + each pair's fares.csv, and writes
/home/z/my-project/download/vexere_nxn_trips.xlsx with:
  - A "Summary" sheet (overall stats + top routes + all 72 pairs table)
  - One sheet per (from, to) pair with trips > 0 (61 sheets)
    Sheet name format: "<from_slug>_to_<to_slug>" (max 31 chars)

Each pair sheet shows the flat fares table (one row per departure) with
key columns: date, time, operator, vehicle, prices (original/discount/final),
seats (available/total), duration.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

# Import the xlsx skill's base templates.
XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, XLSX_SKILL_DIR + "/templates"]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import base as style  # templates/base.py — design tokens & factories

# ---------------------------------------------------------------------------
# City ID → slug mapping (for sheet names; max 31 chars including "_to_").
# ---------------------------------------------------------------------------

CITY_SLUGS = {
    24:  "HaNoi",
    29:  "HCM",
    15:  "DaNang",
    27:  "HaiPhong",
    705: "Hue",
    457: "DaLat",
    482: "SaPa",
    417: "NhaTrang",
    278: "HaGiang",
}

CITY_NAMES = {
    24:  "Hà Nội",
    29:  "Hồ Chí Minh",
    15:  "Đà Nẵng",
    27:  "Hải Phòng",
    705: "Huế",
    457: "Đà Lạt",
    482: "Sa Pa",
    417: "Nha Trang",
    278: "Hà Giang",
}

# ---------------------------------------------------------------------------
# Config.
# ---------------------------------------------------------------------------

MANIFEST_PATH = Path("/tmp/nxn_out/manifest.json")
OUTPUT_PATH = Path("/home/z/my-project/download/vexere_nxn_trips.xlsx")
OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

# Columns to show on each pair sheet (subset of fares.csv columns).
PAIR_SHEET_COLUMNS = [
    "date", "departure_time", "operator_name", "vehicle_type",
    "original_price_vnd", "discount_vnd", "final_price_vnd",
    "available_seats", "total_seats", "duration_minutes",
    "departure_iso", "arrival_iso", "schedule_id",
]

# Friendly column headers (display) for the pair sheets.
COLUMN_HEADERS = {
    "date":                "Date",
    "departure_time":      "Dep. Time",
    "operator_name":       "Operator",
    "vehicle_type":        "Vehicle Type",
    "original_price_vnd":  "Original (VND)",
    "discount_vnd":        "Discount (VND)",
    "final_price_vnd":     "Final (VND)",
    "available_seats":     "Avail. Seats",
    "total_seats":         "Total Seats",
    "duration_minutes":    "Duration (min)",
    "departure_iso":       "Departure ISO",
    "arrival_iso":         "Arrival ISO",
    "schedule_id":         "Schedule ID",
}

# Column widths (in Excel character units).
COLUMN_WIDTHS = {
    "date":                12,
    "departure_time":      10,
    "operator_name":       22,
    "vehicle_type":        22,
    "original_price_vnd":  14,
    "discount_vnd":        13,
    "final_price_vnd":     13,
    "available_seats":     11,
    "total_seats":         11,
    "duration_minutes":    13,
    "departure_iso":       26,
    "arrival_iso":         26,
    "schedule_id":         28,
}


def pair_sheet_name(from_id: int, to_id: int) -> str:
    """Build a sheet name like 'HaNoi_to_SaPa' (max 31 chars)."""
    name = f"{CITY_SLUGS[from_id]}_to_{CITY_SLUGS[to_id]}"
    assert len(name) <= 31, f"sheet name too long: {name}"
    # Excel forbids: \ / ? * [ ] :
    return name


def read_fares_csv(pair_dir: Path) -> pd.DataFrame:
    """Read the fares.csv from a pair directory, returning a DataFrame."""
    fares_path = pair_dir / "fares.csv"
    if not fares_path.exists():
        return pd.DataFrame()
    df = pd.read_csv(fares_path)
    # Sort by date then departure_time for readability.
    if "date" in df.columns and "departure_time" in df.columns:
        df = df.sort_values(["date", "departure_time"]).reset_index(drop=True)
    return df


def style_header_row(ws, row: int, col_start: int, col_end: int):
    """Apply header style to a row range."""
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = style.fill_header()
        cell.font = style.font_header()
        cell.alignment = style.align_header()
        cell.border = style.border_header()
    ws.row_dimensions[row].height = 22


def style_data_cell(cell, row_index: int, is_number: bool = False):
    """Apply alternating-row data style + font."""
    cell.fill = style.fill_data_row(row_index)
    cell.font = style.font_body()
    if is_number:
        cell.alignment = style.align_number()
    else:
        cell.alignment = style.align_text()


def setup_sheet(ws, title: str, last_col: int):
    """Apply standard sheet setup: hide gridlines, set margins, write title."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2  # margin
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 6
    # Title at B2.
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    cell = ws.cell(row=2, column=2, value=title)
    cell.font = Font(name=style.FONT_NAME, size=14, bold=True, color=style.PRIMARY)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def write_pair_sheet(wb: Workbook, rec: dict, df: pd.DataFrame) -> str:
    """Write one pair's fares as a sheet. Returns the sheet name."""
    from_id = rec["from_id"]
    to_id = rec["to_id"]
    sheet_name = pair_sheet_name(from_id, to_id)
    ws = wb.create_sheet(title=sheet_name)

    title = (f"{CITY_NAMES[from_id]} → {CITY_NAMES[to_id]}   "
             f"(area {from_id} → {to_id})   "
             f"{len(df):,} departures across {df['date'].nunique() if 'date' in df else 0} days")

    # Pick columns that actually exist in the DataFrame.
    cols_present = [c for c in PAIR_SHEET_COLUMNS if c in df.columns]
    last_col = 1 + len(cols_present)  # col 1 is margin

    setup_sheet(ws, title, last_col)

    # Column widths.
    for i, col in enumerate(cols_present, start=2):
        ws.column_dimensions[get_column_letter(i)].width = COLUMN_WIDTHS.get(col, 14)

    # Header row at row 4.
    header_row = 4
    for i, col in enumerate(cols_present, start=2):
        cell = ws.cell(row=header_row, column=i, value=COLUMN_HEADERS.get(col, col))
    style_header_row(ws, header_row, 2, last_col)

    # Data rows.
    for r_idx, row in enumerate(df.itertuples(index=False), start=0):
        excel_row = header_row + 1 + r_idx
        is_alt = r_idx % 2 == 1
        for c_idx, col in enumerate(cols_present, start=2):
            val = getattr(row, col, None)
            # Handle NaN → empty string.
            if pd.isna(val):
                val = ""
            cell = ws.cell(row=excel_row, column=c_idx, value=val)
            is_number = col in {
                "original_price_vnd", "discount_vnd", "final_price_vnd",
                "available_seats", "total_seats", "duration_minutes",
            }
            style_data_cell(cell, r_idx, is_number=is_number)
            if is_number and isinstance(val, (int, float)) and val != "":
                if "price_vnd" in col or col == "discount_vnd":
                    cell.number_format = "#,##0"
        ws.row_dimensions[excel_row].height = 16

    # Freeze panes at row 5, col B (so header stays visible on scroll).
    ws.freeze_panes = "B5"

    # Auto-filter on the data range.
    last_row = header_row + len(df)
    ws.auto_filter.ref = f"B{header_row}:{get_column_letter(last_col)}{last_row}"

    return sheet_name


def write_summary_sheet(wb: Workbook, manifest: list[dict]):
    """Write the summary/overview sheet."""
    ws = wb.active
    ws.title = "Summary"

    setup_sheet(ws, "VeXeRe N×N Crawl — Summary", 8)

    # Column widths.
    widths = {2: 22, 3: 14, 4: 14, 5: 14, 6: 22, 7: 18, 8: 18}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    # --- KPI block at row 4-5.
    total_pairs = len(manifest)
    nonzero = [r for r in manifest if r["n_trips"] > 0]
    zero = [r for r in manifest if r["n_trips"] == 0]
    total_trips = sum(r["n_trips"] for r in manifest)
    total_fares = sum(r["n_fares"] for r in manifest)

    kpis = [
        ("Pairs attempted",   f"{total_pairs}"),
        ("Pairs with trips",  f"{len(nonzero)}"),
        ("Pairs (no service)", f"{len(zero)}"),
        ("Total trips",       f"{total_trips:,}"),
        ("Total fares",       f"{total_fares:,}"),
    ]
    for i, (label, val) in enumerate(kpis):
        col = 2 + i
        lbl_cell = ws.cell(row=4, column=col, value=label)
        lbl_cell.font = Font(name=style.FONT_NAME, size=9, color=style.NEUTRAL_600)
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        val_cell = ws.cell(row=5, column=col, value=val)
        val_cell.font = Font(name=style.FONT_NAME, size=14, bold=True, color=style.PRIMARY)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 24

    # --- Top 10 routes by trip count.
    ws.cell(row=7, column=2, value="Top 10 routes by trip count").font = Font(
        name=style.FONT_NAME, size=11, bold=True, color=style.PRIMARY)
    ws.row_dimensions[7].height = 18

    top_headers = ["From", "To", "Trips", "Fares", "Avg Final (VND)"]
    for i, h in enumerate(top_headers, start=2):
        ws.cell(row=8, column=i, value=h)
    style_header_row(ws, 8, 2, 6)

    top10 = sorted(nonzero, key=lambda x: -x["n_trips"])[:10]
    for r_idx, r in enumerate(top10):
        excel_row = 9 + r_idx
        # Compute avg final price from the fares.csv.
        avg_price = ""
        fares_path = Path(r["pair_dir"]) / "fares.csv"
        if fares_path.exists():
            try:
                df = pd.read_csv(fares_path, usecols=["final_price_vnd"])
                if len(df) > 0:
                    avg_price = int(df["final_price_vnd"].mean())
            except Exception:
                pass
        cells = [
            CITY_NAMES.get(r["from_id"], str(r["from_id"])),
            CITY_NAMES.get(r["to_id"],   str(r["to_id"])),
            r["n_trips"],
            r["n_fares"],
            avg_price,
        ]
        for c_idx, val in enumerate(cells, start=2):
            cell = ws.cell(row=excel_row, column=c_idx, value=val)
            is_number = isinstance(val, (int, float)) and not isinstance(val, bool)
            style_data_cell(cell, r_idx, is_number=is_number)
            if c_idx in {4, 5}:  # trips, fares
                cell.number_format = "#,##0"
            elif c_idx == 6 and isinstance(avg_price, int):
                cell.number_format = "#,##0"
        ws.row_dimensions[excel_row].height = 16

    # --- Full pair inventory.
    inv_start = 21
    ws.cell(row=inv_start, column=2, value="All 72 pair combinations").font = Font(
        name=style.FONT_NAME, size=11, bold=True, color=style.PRIMARY)
    ws.row_dimensions[inv_start].height = 18

    inv_headers = ["From ID", "From", "To ID", "To", "Trips", "Fares", "Sheet"]
    for i, h in enumerate(inv_headers, start=2):
        ws.cell(row=inv_start + 1, column=i, value=h)
    style_header_row(ws, inv_start + 1, 2, 8)
    ws.column_dimensions[get_column_letter(8)].width = 24

    # Sort by from_id then to_id.
    sorted_manifest = sorted(manifest, key=lambda r: (r["from_id"], r["to_id"]))
    for r_idx, r in enumerate(sorted_manifest):
        excel_row = inv_start + 2 + r_idx
        sheet_ref = ""
        if r["n_trips"] > 0:
            sheet_ref = pair_sheet_name(r["from_id"], r["to_id"])
        cells = [
            r["from_id"],
            CITY_NAMES.get(r["from_id"], "?"),
            r["to_id"],
            CITY_NAMES.get(r["to_id"], "?"),
            r["n_trips"],
            r["n_fares"],
            sheet_ref,
        ]
        for c_idx, val in enumerate(cells, start=2):
            cell = ws.cell(row=excel_row, column=c_idx, value=val)
            is_number = isinstance(val, (int, float)) and not isinstance(val, bool)
            style_data_cell(cell, r_idx, is_number=is_number)
            if c_idx in {6, 7}:
                cell.number_format = "#,##0"
            if c_idx == 8 and sheet_ref:
                cell.font = Font(name=style.FONT_NAME, size=10, color="0563C1", underline="single")
                cell.hyperlink = f"#'{sheet_ref}'!B2"
        ws.row_dimensions[excel_row].height = 16

    ws.freeze_panes = f"B{inv_start + 2}"


def main() -> int:
    print(f"==> reading manifest: {MANIFEST_PATH}")
    with open(MANIFEST_PATH) as f:
        manifest = json.load(f)

    pairs_with_trips = [r for r in manifest if r["n_trips"] > 0]
    pairs_zero = [r for r in manifest if r["n_trips"] == 0]
    print(f"    {len(manifest)} pairs in manifest")
    print(f"    {len(pairs_with_trips)} pairs with trips > 0")
    print(f"    {len(pairs_zero)} pairs with 0 trips")
    print(f"    {sum(r['n_trips'] for r in manifest):,} total trips")
    print(f"    {sum(r['n_fares'] for r in manifest):,} total fares")

    # Initialize workbook.
    wb = Workbook()
    wb.properties.creator = "Z.ai"
    # Use the default professional palette (deep blue).
    style.use_palette_explicit("professional")

    # Write Summary sheet first (it becomes the active sheet).
    write_summary_sheet(wb, manifest)
    print("==> wrote Summary sheet")

    # Write one sheet per pair with trips > 0, sorted by from_id then to_id.
    pairs_sorted = sorted(pairs_with_trips, key=lambda r: (r["from_id"], r["to_id"]))
    total_rows = 0
    for i, rec in enumerate(pairs_sorted, 1):
        df = read_fares_csv(Path(rec["pair_dir"]))
        if df.empty:
            continue
        sheet_name = write_pair_sheet(wb, rec, df)
        total_rows += len(df)
        if i % 10 == 0 or i == len(pairs_sorted):
            print(f"    [{i}/{len(pairs_sorted)}] wrote {sheet_name} ({len(df):,} rows)")
    print(f"==> wrote {len(pairs_sorted)} pair sheets, {total_rows:,} total rows")

    # Save.
    print(f"==> saving to {OUTPUT_PATH}")
    wb.save(OUTPUT_PATH)
    size_mb = OUTPUT_PATH.stat().st_size / (1024 * 1024)
    print(f"==> done — {size_mb:.2f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
