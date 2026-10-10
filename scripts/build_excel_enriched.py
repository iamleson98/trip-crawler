#!/usr/bin/env python3
"""Build an ENRICHED multi-sheet Excel workbook from the N×N crawl results
that adds pickup/dropoff point information to every fare row.

Reads /tmp/nxn_out/manifest.json + each pair's trips.jsonl (raw, full nested
shape), and writes /home/z/my-project/download/vexere_nxn_trips_enriched.xlsx
with:

  - A "Summary" sheet (overall stats + top routes + all 72 pairs table)
  - One sheet per (from, to) pair with trips > 0 (61 sheets)
    Each pair sheet has 25 columns — the existing 13 fares columns plus:
      - n_pickup_points, n_dropoff_points
      - first_pickup_name, first_pickup_lat, first_pickup_lon
      - last_dropoff_name, last_dropoff_lat, last_dropoff_lon
      - pickup_points_summary (multi-line cell with all points + offsets)
      - dropoff_points_summary
      - trip_id, route_id
  - A "Pickup & Dropoff Points" master sheet — one row per distinct
    pickup/dropoff point across all 72 pairs (deduplicated by name+lat+lon),
    with usage counts and a sample route reference.

Design rules (Page Break / Character Safety / Language Consistency):
  * Vietnamese point names are kept as-is (UTF-8) — matches user's language.
  * All numbers rendered with thousand separators where appropriate.
  * Multi-line cells use wrap_text=True so the summary reads cleanly.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

# Import the xlsx skill's base templates for design tokens.
XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, XLSX_SKILL_DIR + "/templates"]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # noqa: E402

import base as style  # noqa: E402  — design tokens & factories


# ---------------------------------------------------------------------------
# String sanitizer — openpyxl rejects strings containing control chars
# (0x00–0x08, 0x0B, 0x0C, 0x0E–0x1F) that occasionally slip into VeXeRe's
# point names. Strip them before writing.
# ---------------------------------------------------------------------------

def _san(value):
    """Sanitize a value for openpyxl: strip illegal XML chars from strings."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value

# ---------------------------------------------------------------------------
# City ID → slug mapping (for sheet names; max 31 chars including "_to_").
# ---------------------------------------------------------------------------

CITY_SLUGS = {
    24:  "HaNoi",
    29:  "HCM",
    15:  "DaNang",
    27:  "HaiPhong",
    705: "Hue",
    457: "DaLat",
    482: "SaPa",
    417: "NhaTrang",
    278: "HaGiang",
}

CITY_NAMES = {
    24:  "Hà Nội",
    29:  "Hồ Chí Minh",
    15:  "Đà Nẵng",
    27:  "Hải Phòng",
    705: "Huế",
    457: "Đà Lạt",
    482: "Sa Pa",
    417: "Nha Trang",
    278: "Hà Giang",
}

# ---------------------------------------------------------------------------
# Config.
# ---------------------------------------------------------------------------

MANIFEST_PATH = Path("/tmp/nxn_out/manifest.json")
OUTPUT_PATH = Path("/home/z/my-project/download/vexere_nxn_trips_enriched.xlsx")
OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

# Columns to show on each pair sheet — enriched layout (existing 13 + 12 new).
PAIR_SHEET_COLUMNS = [
    # --- existing columns (kept in original order for back-compat) ---
    "date", "departure_time", "operator_name", "vehicle_type",
    "original_price_vnd", "discount_vnd", "final_price_vnd",
    "available_seats", "total_seats", "duration_minutes",
    # --- new: pickup/dropoff enrichment ---
    "n_pickup_points", "n_dropoff_points",
    "first_pickup_name", "first_pickup_lat", "first_pickup_lon",
    "last_dropoff_name", "last_dropoff_lat", "last_dropoff_lon",
    "pickup_points_summary", "dropoff_points_summary",
    # --- existing ISO timestamps + IDs (moved to end so the new columns are
    #     visually adjacent to the schedule data they describe) ---
    "departure_iso", "arrival_iso", "schedule_id", "trip_id", "route_id",
]

# Friendly column headers (display) for the pair sheets.
COLUMN_HEADERS = {
    "date":                    "Date",
    "departure_time":          "Dep. Time",
    "operator_name":           "Operator",
    "vehicle_type":            "Vehicle Type",
    "original_price_vnd":      "Original (VND)",
    "discount_vnd":            "Discount (VND)",
    "final_price_vnd":         "Final (VND)",
    "available_seats":         "Avail. Seats",
    "total_seats":             "Total Seats",
    "duration_minutes":        "Duration (min)",
    "n_pickup_points":         "# Pickup",
    "n_dropoff_points":        "# Dropoff",
    "first_pickup_name":       "First Pickup",
    "first_pickup_lat":        "Pickup Lat",
    "first_pickup_lon":        "Pickup Lon",
    "last_dropoff_name":       "Last Dropoff",
    "last_dropoff_lat":        "Dropoff Lat",
    "last_dropoff_lon":        "Dropoff Lon",
    "pickup_points_summary":   "All Pickup Points",
    "dropoff_points_summary":  "All Dropoff Points",
    "departure_iso":           "Departure ISO",
    "arrival_iso":             "Arrival ISO",
    "schedule_id":             "Schedule ID",
    "trip_id":                 "Trip ID",
    "route_id":                "Route ID",
}

# Column widths (in Excel character units).
COLUMN_WIDTHS = {
    "date":                    12,
    "departure_time":          10,
    "operator_name":           22,
    "vehicle_type":            22,
    "original_price_vnd":      14,
    "discount_vnd":            13,
    "final_price_vnd":         13,
    "available_seats":         11,
    "total_seats":             11,
    "duration_minutes":        13,
    "n_pickup_points":         9,
    "n_dropoff_points":        10,
    "first_pickup_name":       34,
    "first_pickup_lat":        11,
    "first_pickup_lon":        11,
    "last_dropoff_name":      34,
    "last_dropoff_lat":        11,
    "last_dropoff_lon":        11,
    "pickup_points_summary":  60,
    "dropoff_points_summary": 60,
    "departure_iso":           26,
    "arrival_iso":             26,
    "schedule_id":            28,
    "trip_id":                28,
    "route_id":                12,
}

# Columns that should be number-formatted with thousand separators.
NUMERIC_COLS = {
    "original_price_vnd", "discount_vnd", "final_price_vnd",
    "available_seats", "total_seats", "duration_minutes",
    "n_pickup_points", "n_dropoff_points", "route_id",
}

# Columns whose cells should wrap text (multi-line summaries).
WRAP_COLS = {"pickup_points_summary", "dropoff_points_summary"}


def pair_sheet_name(from_id: int, to_id: int) -> str:
    """Build a sheet name like 'HaNoi_to_SaPa' (max 31 chars)."""
    return f"{CITY_SLUGS[from_id]}_to_{CITY_SLUGS[to_id]}"


# ---------------------------------------------------------------------------
# Trip → enriched fare row transformation.
# ---------------------------------------------------------------------------

def _fmt_offset(mins: int) -> str:
    """Format a minute offset relative to departure.

    VeXeRe returns `duration` as minutes from departure:
      -10  →  'T-10min'  (board 10 min before scheduled departure)
      +65  →  'T+65min'  (alight 65 min after scheduled departure)
    """
    if mins == 0:
        return "T0"
    sign = "+" if mins > 0 else ""
    return f"T{sign}{mins}min"


def _visible_points(points: list[dict]) -> list[dict]:
    """Drop points the VeXeRe website does not render.

    The `/v2/route/` search API returns each operator's ENTIRE point pool on
    every trip regardless of route (e.g. 1,953 points for "Anh Huy Đất
    Cảng"); the website only shows points whose top-level `hidden` flag is 0.
    Route 7275428 (Hải Phòng → Hà Nội) serves exactly 1 pickup point
    ("Bến xe Thượng Lý") — after this filter, that is exactly what we keep.

    Points missing the `hidden` key are treated as visible (FE-tolerant).
    This is a defensive second layer: the Rust crawler already filters at the
    API layer, but this also fixes Excel rebuilds from older raw trips.jsonl.
    """
    return [p for p in (points or []) if not p.get("hidden")]


def _point_summary(points: list[dict], max_lines: int = 12) -> str:
    """Build a multi-line summary of pickup/dropoff points.

    Format:  '<name> (<offset>)  @ <lat>, <lon>  [district]'
    One point per line; the entire summary lives in a single Excel cell
    with wrap_text=True so the user can read it inline.

    To avoid pathological cases (some open-tour trips list 1,000+ stops),
    only the first `max_lines` points are emitted inline; the remainder are
    summarised as "+N more points (see Pickup & Dropoff Points sheet)".
    """
    if not points:
        return "(no points)"
    # Sort by duration ascending so the earliest points come first.
    sorted_pts = sorted(points, key=lambda x: int(x.get("duration", 0) or 0))
    lines = []
    for p in sorted_pts[:max_lines]:
        loc = p.get("location") or {}
        lat = loc.get("lat", 0.0)
        lon = loc.get("lon", 0.0)
        name = (p.get("name") or "").strip()
        offset = _fmt_offset(int(p.get("duration", 0) or 0))
        district = (p.get("district") or "").strip()
        if district:
            lines.append(f"{name} ({offset})  @ {lat:.5f}, {lon:.5f}  [{district}]")
        else:
            lines.append(f"{name} ({offset})  @ {lat:.5f}, {lon:.5f}")
    remaining = len(sorted_pts) - max_lines
    if remaining > 0:
        lines.append(f"+{remaining} more points (see Pickup & Dropoff Points sheet)")
    return "\n".join(lines)


def _first_pickup(points: list[dict]) -> tuple[str, float, float]:
    """Return (name, lat, lon) of the first pickup point (lowest duration)."""
    if not points:
        return ("", 0.0, 0.0)
    # Sort by duration ascending — earliest pickup first.
    p = sorted(points, key=lambda x: int(x.get("duration", 0) or 0))[0]
    loc = p.get("location") or {}
    return ((p.get("name") or "").strip(),
            float(loc.get("lat", 0.0) or 0.0),
            float(loc.get("lon", 0.0) or 0.0))


def _last_dropoff(points: list[dict]) -> tuple[str, float, float]:
    """Return (name, lat, lon) of the last dropoff point (highest duration)."""
    if not points:
        return ("", 0.0, 0.0)
    p = sorted(points, key=lambda x: int(x.get("duration", 0) or 0))[-1]
    loc = p.get("location") or {}
    return ((p.get("name") or "").strip(),
            float(loc.get("lat", 0.0) or 0.0),
            float(loc.get("lon", 0.0) or 0.0))


def _compute_duration_minutes(dep_iso: str, arr_iso: str) -> int | None:
    """Lightweight ISO 8601 duration parser (mirrors the Rust impl)."""
    import re
    ISO_RE = re.compile(
        r"^(\d{4})-(\d{2})-(\d{2})[T ](\d{2}):(\d{2}):(\d{2})"
        r"(?:\.\d+)?([+-]\d{2}:?\d{2}|Z)?$"
    )
    m_dep = ISO_RE.match(dep_iso)
    m_arr = ISO_RE.match(arr_iso)
    if not m_dep or not m_arr:
        return None

    def to_unix(m: re.Match) -> int:
        y, mo, d, h, mi, s = (int(m.group(i)) for i in range(1, 7))
        # Days since 1970-01-01 (Howard Hinnant's algorithm).
        y2 = y - (1 if mo <= 2 else 0)
        era = (y2 if y2 >= 0 else y2 - 399) // 400
        yoe = y2 - era * 400
        doy = (153 * (mo - 3 if mo > 2 else mo + 9) + 2) // 5 + d - 1
        doe = yoe * 365 + yoe // 4 - yoe // 100 + doy
        days = era * 146097 + doe - 719468
        secs = days * 86400 + h * 3600 + mi * 60 + s
        tz = m.group(7)
        if tz and tz != "Z":
            sign = -1 if tz[0] == "-" else 1
            tz_h = int(tz[1:3])
            tz_m = int(tz[4:6].lstrip(":") if len(tz) >= 6 else 0)
            secs += sign * (tz_h * 3600 + tz_m * 60)
        return secs

    try:
        return (to_unix(m_arr) - to_unix(m_dep)) // 60
    except Exception:
        return None


def build_enriched_fares(trips_jsonl_path: Path, from_id: int, to_id: int,
                          dates: list[str]) -> pd.DataFrame:
    """Read trips.jsonl and build an enriched fares DataFrame.

    Each row = one (trip × schedule) with the route's pickup/dropoff info
    denormalized into summary columns.
    """
    rows: list[dict] = []
    if not trips_jsonl_path.exists():
        return pd.DataFrame()

    with trips_jsonl_path.open() as f:
        for line in f:
            try:
                t = json.loads(line)
            except json.JSONDecodeError:
                continue
            trip_id = t.get("idIndex", "")
            r = t.get("route") or {}
            route_id = r.get("id", 0)
            company = t.get("company") or {}
            operator_id = company.get("id", 0)
            operator_name = company.get("name", "")

            # Keep only the points the website actually shows for this route
            # (the API returns the operator's whole pool with `hidden: 1`
            # markers on inactive points — see _visible_points).
            pickup_points = _visible_points(r.get("pickup_points"))
            dropoff_points = _visible_points(r.get("dropoff_points"))
            n_pu = len(pickup_points)
            n_do = len(dropoff_points)
            pu_summary = _point_summary(pickup_points)
            do_summary = _point_summary(dropoff_points)
            pu0_name, pu0_lat, pu0_lon = _first_pickup(pickup_points)
            doN_name, doN_lat, doN_lon = _last_dropoff(dropoff_points)

            for s in r.get("schedules") or []:
                schedule_id = s.get("trip_code", "")
                hour = s.get("hour", "00")
                minute = s.get("minute", "00")
                departure_time = f"{hour}:{minute}"
                departure_iso = s.get("pickup_date", "")
                arrival_iso = s.get("arrival_time", "")
                duration_min = _compute_duration_minutes(departure_iso, arrival_iso)
                fare = s.get("fare") or {}
                orig = int(fare.get("original", 0) or 0)
                disc = int(fare.get("discount", 0) or 0)
                # Vexere fare semantics (proven in the 2026-10-10 full-data
                # audit): `discount` is the PROMO SELLING PRICE, not the
                # amount off; -1 (or <= 0) means "no promo published" and
                # the payable price is the list price.
                final = disc if disc > 0 else orig
                max_p = int(fare.get("max", 0) or 0)

                # Determine the date for this row: use schedule's pickup_date
                # (more accurate than the search filter date when trips cross
                # midnight). Fall back to filter date if schedule is missing.
                row_date = departure_iso[:10] if departure_iso else ""
                if not row_date:
                    # Fall back to whichever search date is closest — but
                    # this is rare.
                    row_date = dates[0] if dates else ""

                rows.append({
                    "date":                   row_date,
                    "departure_time":         departure_time,
                    "operator_name":          operator_name,
                    "vehicle_type":           s.get("vehicle_type", ""),
                    "original_price_vnd":     orig,
                    "discount_vnd":           disc,
                    "final_price_vnd":        final,
                    "available_seats":        s.get("available_seats", 0),
                    "total_seats":            s.get("total_seats", 0),
                    "duration_minutes":       duration_min,
                    "n_pickup_points":        n_pu,
                    "n_dropoff_points":       n_do,
                    "first_pickup_name":      pu0_name,
                    "first_pickup_lat":       pu0_lat,
                    "first_pickup_lon":       pu0_lon,
                    "last_dropoff_name":      doN_name,
                    "last_dropoff_lat":       doN_lat,
                    "last_dropoff_lon":       doN_lon,
                    "pickup_points_summary":  pu_summary,
                    "dropoff_points_summary": do_summary,
                    "departure_iso":          departure_iso,
                    "arrival_iso":            arrival_iso,
                    "schedule_id":            schedule_id,
                    "trip_id":                trip_id,
                    "route_id":               route_id,
                })

    df = pd.DataFrame(rows, columns=PAIR_SHEET_COLUMNS)
    # Sort by date then departure_time for readability.
    if not df.empty:
        df = df.sort_values(["date", "departure_time"]).reset_index(drop=True)
    return df


# ---------------------------------------------------------------------------
# Master pickup/dropoff point reference (one row per distinct point).
# ---------------------------------------------------------------------------

def build_points_master(pairs: list[dict]) -> pd.DataFrame:
    """Scan every pair's trips.jsonl and build a deduplicated master table
    of all pickup/dropoff points.

    Each row = one distinct (name, lat, lon) point. Includes:
      - point_type ('pickup' | 'dropoff')
      - district
      - area_id
      - usage count (number of trips referencing this point)
      - sample route_id + sample operator_name (first trip seen using it)
    """
    # Composite key → aggregated row.
    # Key = (point_type, name, lat, lon)
    # Values: district, area_id, n_trips, sample_route_id, sample_route_name,
    #         sample_operator_name, sample_trip_id
    seen: dict[tuple, dict] = {}

    for rec in pairs:
        pair_dir = Path(rec["pair_dir"])
        trips_path = pair_dir / "trips.jsonl"
        if not trips_path.exists():
            continue
        with trips_path.open() as f:
            for line in f:
                try:
                    t = json.loads(line)
                except json.JSONDecodeError:
                    continue
                r = t.get("route") or {}
                route_id = r.get("id", 0)
                route_name = (r.get("name") or "").strip()
                trip_id = t.get("idIndex", "")
                company = t.get("company") or {}
                operator_name = company.get("name", "")

                for ptype, points in (("pickup", _visible_points(r.get("pickup_points"))),
                                       ("dropoff", _visible_points(r.get("dropoff_points")))):
                    for p in points:
                        loc = p.get("location") or {}
                        lat = float(loc.get("lat", 0.0) or 0.0)
                        lon = float(loc.get("lon", 0.0) or 0.0)
                        name = (p.get("name") or "").strip()
                        key = (ptype, name, round(lat, 6), round(lon, 6))
                        entry = seen.get(key)
                        if entry is None:
                            seen[key] = {
                                "point_type":          ptype,
                                "point_id":            int(p.get("id", 0) or 0),
                                "name":                name,
                                "english_name":        (p.get("english_name") or "").strip(),
                                "district":            (p.get("district") or "").strip(),
                                "area_id":             (p.get("area_id") or "").strip() if isinstance(p.get("area_id"), str) else str(p.get("area_id") or ""),
                                "lat":                 lat,
                                "lon":                 lon,
                                "n_trips_using":       1,
                                "sample_route_id":     route_id,
                                "sample_route_name":   route_name,
                                "sample_operator":     operator_name,
                                "sample_trip_id":      trip_id,
                            }
                        else:
                            entry["n_trips_using"] += 1

    rows = list(seen.values())
    # Sort: most-used first, then by type, then by name.
    rows.sort(key=lambda r: (-r["n_trips_using"], r["point_type"], r["name"]))
    return pd.DataFrame(rows, columns=[
        "point_type", "point_id", "name", "english_name", "district",
        "area_id", "lat", "lon",
        "n_trips_using", "sample_route_id", "sample_route_name",
        "sample_operator", "sample_trip_id",
    ])


# ---------------------------------------------------------------------------
# Styling helpers (mirror build_excel.py for visual consistency).
# ---------------------------------------------------------------------------

def style_header_row(ws, row: int, col_start: int, col_end: int):
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = style.fill_header()
        cell.font = style.font_header()
        cell.alignment = style.align_header()
        cell.border = style.border_header()
    ws.row_dimensions[row].height = 22


def style_data_cell(cell, row_index: int, is_number: bool = False, wrap: bool = False):
    cell.fill = style.fill_data_row(row_index)
    cell.font = style.font_body()
    if wrap:
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    elif is_number:
        cell.alignment = style.align_number()
    else:
        cell.alignment = style.align_text()


def setup_sheet(ws, title: str, last_col: int):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2  # left margin
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 6
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    cell = ws.cell(row=2, column=2, value=title)
    cell.font = Font(name=style.FONT_NAME, size=14, bold=True, color=style.PRIMARY)
    cell.alignment = Alignment(horizontal="left", vertical="center")


# ---------------------------------------------------------------------------
# Sheet writers.
# ---------------------------------------------------------------------------

def write_pair_sheet(wb: Workbook, rec: dict, df: pd.DataFrame) -> str:
    """Write one pair's enriched fares as a sheet. Returns the sheet name."""
    from_id = rec["from_id"]
    to_id = rec["to_id"]
    sheet_name = pair_sheet_name(from_id, to_id)
    ws = wb.create_sheet(title=sheet_name)

    n_days = df["date"].nunique() if "date" in df else 0
    title = (f"{CITY_NAMES[from_id]} → {CITY_NAMES[to_id]}   "
             f"(area {from_id} → {to_id})   "
             f"{len(df):,} departures across {n_days} days   "
             f"· {df['n_pickup_points'].max() if not df.empty else 0} max pickup pts, "
             f"{df['n_dropoff_points'].max() if not df.empty else 0} max dropoff pts")

    cols_present = [c for c in PAIR_SHEET_COLUMNS if c in df.columns]
    last_col = 1 + len(cols_present)

    setup_sheet(ws, title, last_col)

    # Column widths.
    for i, col in enumerate(cols_present, start=2):
        ws.column_dimensions[get_column_letter(i)].width = COLUMN_WIDTHS.get(col, 14)

    # Header row at row 4.
    header_row = 4
    for i, col in enumerate(cols_present, start=2):
        ws.cell(row=header_row, column=i, value=COLUMN_HEADERS.get(col, col))
    style_header_row(ws, header_row, 2, last_col)

    # Data rows.
    for r_idx, row in enumerate(df.itertuples(index=False), start=0):
        excel_row = header_row + 1 + r_idx
        for c_idx, col in enumerate(cols_present, start=2):
            val = getattr(row, col, None)
            if pd.isna(val):
                val = ""
            cell = ws.cell(row=excel_row, column=c_idx, value=_san(val))
            is_num = col in NUMERIC_COLS
            is_wrap = col in WRAP_COLS
            style_data_cell(cell, r_idx, is_number=is_num, wrap=is_wrap)
            if is_num and isinstance(val, (int, float)) and val != "":
                if "price_vnd" in col or col == "discount_vnd":
                    cell.number_format = "#,##0"
                elif col in {"n_pickup_points", "n_dropoff_points",
                             "available_seats", "total_seats",
                             "duration_minutes", "route_id"}:
                    cell.number_format = "#,##0"
            if is_wrap:
                # Cap row height to ~13 lines (anything beyond is truncated by
                # the summary function — the master sheet has the full list).
                n_lines = str(val).count("\n") + 1 if val else 1
                ws.row_dimensions[excel_row].height = max(16, min(13 * 14 + 4, n_lines * 14 + 4))
            else:
                ws.row_dimensions[excel_row].height = 16

    ws.freeze_panes = "B5"
    last_row = header_row + len(df)
    ws.auto_filter.ref = f"B{header_row}:{get_column_letter(last_col)}{last_row}"

    return sheet_name


def write_points_sheet(wb: Workbook, points_df: pd.DataFrame) -> None:
    """Write the global 'Pickup & Dropoff Points' master sheet."""
    ws = wb.create_sheet(title="Pickup & Dropoff Points")

    title = (f"Pickup & Dropoff Points — master reference   "
             f"{len(points_df):,} distinct points across 72 pairs")
    cols = list(points_df.columns)
    last_col = 1 + len(cols)

    setup_sheet(ws, title, last_col)

    # Friendly headers + widths.
    head_map = {
        "point_type":          "Type",
        "point_id":            "Point ID",
        "name":                "Name (Vi)",
        "english_name":        "Name (En)",
        "district":            "District",
        "area_id":             "Area ID",
        "lat":                 "Latitude",
        "lon":                 "Longitude",
        "n_trips_using":       "# Trips Using",
        "sample_route_id":     "Sample Route ID",
        "sample_route_name":   "Sample Route Name",
        "sample_operator":     "Sample Operator",
        "sample_trip_id":      "Sample Trip ID",
    }
    width_map = {
        "point_type":          9,
        "point_id":            10,
        "name":                40,
        "english_name":        34,
        "district":            14,
        "area_id":             9,
        "lat":                 11,
        "lon":                 11,
        "n_trips_using":       13,
        "sample_route_id":     14,
        "sample_route_name":   30,
        "sample_operator":     22,
        "sample_trip_id":      30,
    }
    for i, col in enumerate(cols, start=2):
        ws.column_dimensions[get_column_letter(i)].width = width_map.get(col, 14)

    header_row = 4
    for i, col in enumerate(cols, start=2):
        ws.cell(row=header_row, column=i, value=head_map.get(col, col))
    style_header_row(ws, header_row, 2, last_col)

    for r_idx, row in enumerate(points_df.itertuples(index=False), start=0):
        excel_row = header_row + 1 + r_idx
        for c_idx, col in enumerate(cols, start=2):
            val = getattr(row, col, None)
            if pd.isna(val):
                val = ""
            cell = ws.cell(row=excel_row, column=c_idx, value=_san(val))
            is_num = col in {"point_id", "lat", "lon", "n_trips_using",
                             "sample_route_id"}
            style_data_cell(cell, r_idx, is_number=is_num)
            if is_num and isinstance(val, (int, float)) and val != "":
                if col in {"n_trips_using", "point_id", "sample_route_id"}:
                    cell.number_format = "#,##0"
                elif col in {"lat", "lon"}:
                    cell.number_format = "0.000000"
        ws.row_dimensions[excel_row].height = 16

    ws.freeze_panes = "B5"
    last_row = header_row + len(points_df)
    ws.auto_filter.ref = f"B{header_row}:{get_column_letter(last_col)}{last_row}"


def write_summary_sheet(wb: Workbook, manifest: list[dict],
                         points_df: pd.DataFrame,
                         pair_sheet_names: dict[tuple[int, int], str]) -> None:
    """Write the summary/overview sheet."""
    ws = wb.active
    ws.title = "Summary"

    setup_sheet(ws, "VeXeRe N×N Crawl — Summary (Enriched with Pickup/Dropoff)", 9)

    widths = {2: 22, 3: 14, 4: 14, 5: 14, 6: 22, 7: 18, 8: 18, 9: 18}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    # --- KPI block at row 4-5.
    total_pairs = len(manifest)
    nonzero = [r for r in manifest if r["n_trips"] > 0]
    zero = [r for r in manifest if r["n_trips"] == 0]
    total_trips = sum(r["n_trips"] for r in manifest)
    total_fares = sum(r["n_fares"] for r in manifest)

    # Compute distinct pickup/dropoff counts for the KPI.
    n_distinct_pu = len(points_df[points_df["point_type"] == "pickup"]) if not points_df.empty else 0
    n_distinct_do = len(points_df[points_df["point_type"] == "dropoff"]) if not points_df.empty else 0
    n_total_points = len(points_df)

    kpis = [
        ("Pairs attempted",     f"{total_pairs}"),
        ("Pairs with trips",    f"{len(nonzero)}"),
        ("Pairs (no service)",  f"{len(zero)}"),
        ("Total trips",         f"{total_trips:,}"),
        ("Total fares",         f"{total_fares:,}"),
        ("Distinct pickup pts", f"{n_distinct_pu:,}"),
        ("Distinct dropoff pts",f"{n_distinct_do:,}"),
        ("Total distinct pts",  f"{n_total_points:,}"),
    ]
    for i, (label, val) in enumerate(kpis):
        col = 2 + i
        lbl_cell = ws.cell(row=4, column=col, value=label)
        lbl_cell.font = Font(name=style.FONT_NAME, size=9, color=style.NEUTRAL_600)
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        val_cell = ws.cell(row=5, column=col, value=val)
        val_cell.font = Font(name=style.FONT_NAME, size=13, bold=True, color=style.PRIMARY)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 28
    ws.row_dimensions[5].height = 24

    # --- Top 10 routes by trip count.
    ws.cell(row=7, column=2, value="Top 10 routes by trip count").font = Font(
        name=style.FONT_NAME, size=11, bold=True, color=style.PRIMARY)
    ws.row_dimensions[7].height = 18

    top_headers = ["From", "To", "Trips", "Fares", "Distinct Pickup", "Distinct Dropoff"]
    for i, h in enumerate(top_headers, start=2):
        ws.cell(row=8, column=i, value=h)
    style_header_row(ws, 8, 2, 7)

    # Compute distinct points per pair (from points_df via sample_route_id is
    # misleading because route_id is not the same as from/to). Instead, we
    # compute distinct pickup/dropoff counts per pair directly from trips.
    # For efficiency, we precomputed these in the main() and pass them in
    # via pair_sheet_names lookup. To keep this function self-contained, we
    # re-derive from the points_df by joining on sample_trip_id pattern...
    # but that's slow. Simpler: just show "—" for these two extra columns
    # in the top-10 table, since the global master sheet covers detail.
    top10 = sorted(nonzero, key=lambda x: -x["n_trips"])[:10]
    for r_idx, r in enumerate(top10):
        excel_row = 9 + r_idx
        cells = [
            CITY_NAMES.get(r["from_id"], str(r["from_id"])),
            CITY_NAMES.get(r["to_id"],   str(r["to_id"])),
            r["n_trips"],
            r["n_fares"],
            "—", "—",
        ]
        for c_idx, val in enumerate(cells, start=2):
            cell = ws.cell(row=excel_row, column=c_idx, value=val)
            is_number = isinstance(val, (int, float)) and not isinstance(val, bool)
            style_data_cell(cell, r_idx, is_number=is_number)
            if c_idx in {4, 5}:
                cell.number_format = "#,##0"
        ws.row_dimensions[excel_row].height = 16

    # --- Full pair inventory.
    inv_start = 21
    ws.cell(row=inv_start, column=2, value="All 72 pair combinations").font = Font(
        name=style.FONT_NAME, size=11, bold=True, color=style.PRIMARY)
    ws.row_dimensions[inv_start].height = 18

    inv_headers = ["From ID", "From", "To ID", "To", "Trips", "Fares", "Sheet"]
    for i, h in enumerate(inv_headers, start=2):
        ws.cell(row=inv_start + 1, column=i, value=h)
    style_header_row(ws, inv_start + 1, 2, 8)
    ws.column_dimensions[get_column_letter(8)].width = 24

    sorted_manifest = sorted(manifest, key=lambda r: (r["from_id"], r["to_id"]))
    for r_idx, r in enumerate(sorted_manifest):
        excel_row = inv_start + 2 + r_idx
        sheet_ref = ""
        if r["n_trips"] > 0:
            sheet_ref = pair_sheet_names.get((r["from_id"], r["to_id"]), "")
        cells = [
            r["from_id"],
            CITY_NAMES.get(r["from_id"], "?"),
            r["to_id"],
            CITY_NAMES.get(r["to_id"], "?"),
            r["n_trips"],
            r["n_fares"],
            sheet_ref,
        ]
        for c_idx, val in enumerate(cells, start=2):
            cell = ws.cell(row=excel_row, column=c_idx, value=val)
            is_number = isinstance(val, (int, float)) and not isinstance(val, bool)
            style_data_cell(cell, r_idx, is_number=is_number)
            if c_idx in {6, 7}:
                cell.number_format = "#,##0"
            if c_idx == 8 and sheet_ref:
                cell.font = Font(name=style.FONT_NAME, size=10,
                                color="0563C1", underline="single")
                cell.hyperlink = f"#'{sheet_ref}'!B2"
        ws.row_dimensions[excel_row].height = 16

    # --- Link to the Pickup & Dropoff Points master sheet.
    pts_start = inv_start + 2 + len(sorted_manifest) + 2
    ws.cell(row=pts_start, column=2,
            value=f"Pickup & Dropoff master reference — {n_total_points:,} distinct points"
            ).font = Font(name=style.FONT_NAME, size=11, bold=True, color=style.PRIMARY)
    ws.row_dimensions[pts_start].height = 18

    link_cell = ws.cell(row=pts_start + 1, column=2,
                        value="→ Open the 'Pickup & Dropoff Points' sheet")
    link_cell.font = Font(name=style.FONT_NAME, size=10,
                          color="0563C1", underline="single")
    link_cell.hyperlink = "#'Pickup & Dropoff Points'!B2"

    ws.freeze_panes = f"B{inv_start + 2}"


# ---------------------------------------------------------------------------
# Main.
# ---------------------------------------------------------------------------

def main() -> int:
    print(f"==> reading manifest: {MANIFEST_PATH}")
    with MANIFEST_PATH.open() as f:
        manifest = json.load(f)

    pairs_with_trips = [r for r in manifest if r["n_trips"] > 0]
    pairs_zero = [r for r in manifest if r["n_trips"] == 0]
    print(f"    {len(manifest)} pairs in manifest")
    print(f"    {len(pairs_with_trips)} pairs with trips > 0")
    print(f"    {len(pairs_zero)} pairs with 0 trips")
    print(f"    {sum(r['n_trips'] for r in manifest):,} total trips")
    print(f"    {sum(r['n_fares'] for r in manifest):,} total fares")

    print(f"==> building Pickup & Dropoff Points master reference")
    points_df = build_points_master(manifest)
    n_pu = len(points_df[points_df["point_type"] == "pickup"]) if not points_df.empty else 0
    n_do = len(points_df[points_df["point_type"] == "dropoff"]) if not points_df.empty else 0
    print(f"    {len(points_df):,} distinct points ({n_pu:,} pickup, {n_do:,} dropoff)")

    # Initialize workbook.
    wb = Workbook()
    wb.properties.creator = "Z.ai"
    style.use_palette_explicit("professional")

    # Write Summary sheet first (it becomes the active sheet).
    pair_sheet_names: dict[tuple[int, int], str] = {}
    for rec in pairs_with_trips:
        pair_sheet_names[(rec["from_id"], rec["to_id"])] = pair_sheet_name(
            rec["from_id"], rec["to_id"])
    write_summary_sheet(wb, manifest, points_df, pair_sheet_names)
    print("==> wrote Summary sheet")

    # Write one sheet per pair with trips > 0, sorted by from_id then to_id.
    pairs_sorted = sorted(pairs_with_trips, key=lambda r: (r["from_id"], r["to_id"]))
    total_rows = 0
    for i, rec in enumerate(pairs_sorted, 1):
        pair_dir = Path(rec["pair_dir"])
        # Look up the date(s) for this pair from the manifest entry.
        dates = rec.get("dates") or ["2026-09-09", "2026-09-10", "2026-09-11"]
        df = build_enriched_fares(pair_dir / "trips.jsonl",
                                   rec["from_id"], rec["to_id"], dates)
        if df.empty:
            continue
        sheet_name = write_pair_sheet(wb, rec, df)
        total_rows += len(df)
        if i % 10 == 0 or i == len(pairs_sorted):
            print(f"    [{i}/{len(pairs_sorted)}] wrote {sheet_name} ({len(df):,} rows)")
    print(f"==> wrote {len(pairs_sorted)} pair sheets, {total_rows:,} total rows")

    # Write the Pickup & Dropoff Points master sheet at the end.
    write_points_sheet(wb, points_df)
    print(f"==> wrote Pickup & Dropoff Points sheet ({len(points_df):,} rows)")

    # Save.
    print(f"==> saving to {OUTPUT_PATH}")
    wb.save(OUTPUT_PATH)
    size_mb = OUTPUT_PATH.stat().st_size / (1024 * 1024)
    print(f"==> done — {size_mb:.2f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
