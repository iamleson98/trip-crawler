#!/usr/bin/env python3
"""Extract the unique operator (brand) list across all route sheets of the
enriched N×N Excel workbook, with per-brand stats for the research report."""
import json
import openpyxl
from collections import defaultdict

SRC = "/home/z/my-project/download/vexere_nxn_trips_enriched.xlsx"

wb = openpyxl.load_workbook(SRC, read_only=True)
route_sheets = [s for s in wb.sheetnames if s not in ("Summary", "Pickup & Dropoff Points")]

brands = defaultdict(lambda: {"n_trips": 0, "routes": set(), "prices": [], "veh_types": set()})
for sn in route_sheets:
    ws = wb[sn]
    header = None
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        if row and "Operator" in [c for c in row if c]:
            header = list(row)
            break
    if not header:
        continue
    op_i = header.index("Operator")
    # locate other useful columns dynamically
    def col(name):
        return header.index(name) if name in header else None
    date_i, vt_i = col("Date"), col("Vehicle Type")
    fin_i, orig_i = col("Final (VND)"), col("Original (VND)")
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row or op_i >= len(row) or row[op_i] is None:
            continue
        name = str(row[op_i]).strip()
        if not name or name == "Operator":
            continue
        b = brands[name]
        b["n_trips"] += 1
        b["routes"].add(sn.replace("_to_", " → "))
        if vt_i and vt_i < len(row) and row[vt_i]:
            b["veh_types"].add(str(row[vt_i]))
        if fin_i and fin_i < len(row) and isinstance(row[fin_i], (int, float)):
            b["prices"].append(row[fin_i])

out = []
for name, b in brands.items():
    out.append({
        "name": name,
        "n_trips": b["n_trips"],
        "n_routes": len(b["routes"]),
        "routes": sorted(b["routes"]),
        "veh_types": sorted(b["veh_types"]),
        "min_price": min(b["prices"]) if b["prices"] else None,
        "max_price": max(b["prices"]) if b["prices"] else None,
    })
out.sort(key=lambda x: -x["n_trips"])

with open("/home/z/my-project/research/brand_inventory.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)

print(f"UNIQUE BRANDS: {len(out)}")
print(f"TOTAL TRIPS: {sum(x['n_trips'] for x in out)}")
print("\nTOP 30 BY TRIP COUNT:")
for b in out[:30]:
    print(f"  {b['n_trips']:6d} trips | {b['n_routes']:2d} routes | {b['name']}")
print("\nBRANDS WITH >=10 TRIPS: " + str(sum(1 for x in out if x["n_trips"] >= 10)))
print("BRANDS WITH <10 TRIPS:  " + str(sum(1 for x in out if x["n_trips"] < 10)))
