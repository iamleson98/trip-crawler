#!/usr/bin/env python3
"""Build the Vexere research report workbook:
  Summary | Vexere Features | Brand Directory | Support Phones |
  Customer Feedback | Ratings Analysis | Review
Design: xlsx skill design system (templates/base.py, professional palette).
"""
import json
import sys
import os

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (  # noqa: E402
    FONT_NAME, HEADER_BOLD, PRIMARY, NEUTRAL_600, NEUTRAL_900,
    CHART_COLORS, ROW_HEIGHTS,
    font_body, font_caption, font_subheader, font_title,
    setup_sheet, style_header_row, style_data_row, style_total_row,
    auto_fit_columns, create_bar_chart, setup_chart_titles,
    apply_chart_colors, align_number, align_text, align_date,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from openpyxl.chart import Reference  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

OUT = "/home/z/my-project/download/vexere_brands_feedback_report.xlsx"
BRANDS = json.load(open("/home/z/my-project/research/final_brands.json", encoding="utf-8"))
REVIEWS = json.load(open("/home/z/my-project/research/final_reviews.json", encoding="utf-8"))
FEAT = json.load(open("/home/z/my-project/research/final_features.json", encoding="utf-8"))
SUPPORT = FEAT["support"]
FEATURES = FEAT["features"]

DATA_DATE = "2026-09-09"


def safe(v):
    """Guard text that Excel could misread (leading =, +, @)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "@"):
        return "'" + v
    return v


def write_table(ws, headers, rows, start_row=4, number_cols=None, date_cols=None,
                col_start=2, totals=None):
    """Write header + data rows with design-system styling.
    number_cols/date_cols: set of 0-based indices within `headers`."""
    number_cols = number_cols or set()
    date_cols = date_cols or set()
    last_col = col_start + len(headers) - 1
    for i, h in enumerate(headers):
        ws.cell(row=start_row, column=col_start + i, value=h)
    style_header_row(ws, start_row, col_start, last_col)
    r = start_row + 1
    for idx, row in enumerate(rows):
        for i, v in enumerate(row):
            c = ws.cell(row=r, column=col_start + i, value=safe(v))
            if i in number_cols and isinstance(v, (int, float)):
                c.number_format = "#,##0" if isinstance(v, int) else "#,##0.0"
        style_data_row(ws, r, col_start, last_col, idx)
        for i in range(len(headers)):
            cell = ws.cell(row=r, column=col_start + i)
            if i in number_cols:
                cell.alignment = align_number()
            elif i in date_cols:
                cell.alignment = align_date()
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        r += 1
    if totals:
        for i, v in enumerate(totals):
            c = ws.cell(row=r, column=col_start + i, value=v)
            if i in number_cols and isinstance(v, (int, float)):
                c.number_format = "#,##0" if isinstance(v, int) else "#,##0.0"
        style_total_row(ws, r, col_start, last_col)
        r += 1
    return r  # next free row


def caption(ws, row, col, text, span=None):
    c = ws.cell(row=row, column=col, value=safe(text))
    c.font = font_caption()
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    if span:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    return row + 1


def hyperlink(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    if text:
        c.hyperlink = text
        c.font = Font(name=FONT_NAME, size=10, color=PRIMARY, underline="single")
        c.alignment = Alignment(horizontal="left", vertical="top")


wb = Workbook()
wb.properties.creator = "Z.ai"

# ============================================================ 1. SUMMARY
ws = wb.active
ws.title = "Summary"
setup_sheet(ws, title="Vexere.com Platform Research Report — Features, Brand Directory, Support Phones & Customer Feedback", last_col=8)
ws.cell(row=3, column=2, value=f"Research date: {DATA_DATE}  ·  Source: vexere.com brand pages (sitemap), live site & app stores, web search  ·  Companion to the N×N trips dataset (72 route pairs, 41,993 trips)").font = font_caption()
ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=8)
ws.row_dimensions[3].height = 14

n_brands = len(BRANDS)
n_rating = sum(1 for b in BRANDS if b["rating"])
rated = [b["rating"] for b in BRANDS if b["rating"]]
avg_rating = sum(rated) / len(rated) if rated else 0
total_cnt = sum(b["rating_count"] or 0 for b in BRANDS)
n_own = sum(1 for b in BRANDS if b["own_phones"])
n_web = sum(1 for b in BRANDS if b["web_phones"])
n_url = sum(1 for b in BRANDS if b["url"])

metrics = [
    ("Brands covered (from crawled dataset)", n_brands, "Every operator found in the 72-pair N×N crawl (2026-09-09 → 09-11)"),
    ("Brands with a vexere.com page", n_url, "Matched via bus-operator-details sitemap (name ↔ slug)"),
    ("Brands with an average rating", n_rating, "Rating (0–5) shown on vexere brand pages"),
    ("Average brand rating (of rated)", round(avg_rating, 2), "Simple mean across rated brands"),
    ("Cumulative customer ratings", total_cnt, "Sum of rating_count across brands — scale of feedback on vexere"),
    ("Individual reviews captured verbatim", len(REVIEWS), "Latest customer reviews per brand (JSON-LD review[])"),
    ("Brands with operator's own phone(s)", n_own, "'Số điện thoại xe …' blocks on vexere brand pages"),
    ("Brands with web-researched hotlines", n_web, "Official hotlines verified via web search (top brands)"),
    ("Vexere platform features documented", len(FEATURES), "Ticketing, platform, B2B, content, company — see 'Vexere Features'"),
]
nxt = write_table(ws, ["Metric", "Value", "Notes"], metrics, number_cols={1})

nxt += 1
c = ws.cell(row=nxt, column=2, value="How to read this report")
c.font = font_subheader()
nxt += 1
for line in [
    "1. 'Vexere Features' — everything vexere.com offers beyond bus search: trains, flights, motorbike rental, charter, BMS operator software, affiliate network, promotions, payments, apps.",
    "2. 'Brand Directory' — all 306 operators with vexere rating, rating count, price range, booking hotline, operator's own phones, address and page URL.",
    "3. 'Support Phones' — deep-researched support hotlines for the 30 largest brands (web search cross-checked against vexere pages).",
    "4. 'Customer Feedback' — 1,199 verbatim customer reviews collected from vexere.com brand pages (author, date, stars, text).",
    "5. 'Ratings Analysis' — rating distribution, most-reviewed and best-rated brands, star mix of captured reviews, with charts.",
    "6. 'Review' — live cross-check formulas verifying sheet row counts and totals (all should show PASS).",
]:
    nxt = caption(ws, nxt, 2, line, span=7)
    ws.row_dimensions[nxt - 1].height = 15

nxt += 1
c = ws.cell(row=nxt, column=2, value="Primary data sources")
c.font = font_subheader()
nxt += 1
for s in [
    "vexere.com sitemaps: bus-operator-details.xml (~4,300 brand URLs: danh-gia / so-dien-thoai-dia-chi / hinh-anh pages) · crawled 2026-09-09",
    "vexere.com brand pages — JSON-LD (LocalBusiness, Product+aggregateRating+review) + visible contact blocks",
    "Web search (44 queries): official operator websites (futabus.vn, xedienlinh.com, tankimchi.vn, …), App Store / Google Play, blog.vexere.com, bms.vexere.com, press coverage",
    "Crawled trips dataset: vexere_nxn_trips_enriched.xlsx (41,993 trips / 72 pairs / 2026-09-09→11)",
]:
    nxt = caption(ws, nxt, 2, s, span=7)
    ws.row_dimensions[nxt - 1].height = 15

auto_fit_columns(ws, min_width=10, max_width=60, header_row=4, data_start_row=5)
ws.column_dimensions["D"].width = 70
ws.freeze_panes = "B5"

# ============================================================ 2. VEXERE FEATURES
ws = wb.create_sheet("Vexere Features")
setup_sheet(ws, title="Everything vexere.com Offers — Platform Feature Inventory", last_col=6)
ws.cell(row=3, column=2, value=f"Researched {DATA_DATE} via web search, site sitemaps and app-store listings · 5 categories").font = font_caption()
ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=6)

rows = [list(f) for f in FEATURES]
nxt = write_table(ws, ["Category", "Feature", "Details", "Where on vexere.com", "Source"], rows)
nxt += 1
nxt = caption(ws, nxt, 2, f"Vexere support: {SUPPORT['vexere_hotline']} ({SUPPORT['vexere_hours']}) · marketplace CS {SUPPORT['vexere_marketplace_cs']} · {SUPPORT['vexere_email']} · motorbike-rental hotline {SUPPORT['motorbike_hotline']}", span=5)
auto_fit_columns(ws, min_width=12, max_width=30, header_row=4, data_start_row=5)
ws.column_dimensions["D"].width = 78
ws.column_dimensions["C"].width = 26
ws.column_dimensions["E"].width = 26
ws.column_dimensions["F"].width = 26
ws.freeze_panes = "B5"

# ============================================================ 3. BRAND DIRECTORY
ws = wb.create_sheet("Brand Directory")
setup_sheet(ws, title="Brand Directory — 306 Operators on vexere.com", last_col=15)
ws.cell(row=3, column=2, value="Sorted by trips in the crawled dataset · ratings & phones from each brand's vexere.com page · prices in VND").font = font_caption()
ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=15)

hdr = ["#", "Brand", "Rating", "# Ratings", "Trips (dataset)", "# Routes", "Main routes", "Vehicle types",
       "Price range (vexere)", "Vexere booking hotline", "Operator's own phones", "Address", "Vexere page", "Notes"]
rows = []
for i, b in enumerate(BRANDS, 1):
    price = b["price_range"] or (f"{b['min_price']:,} – {b['max_price']:,} VND" if b["min_price"] is not None else "")
    rows.append([
        i, b["name"], b["rating"] if b["rating"] else None, b["rating_count"] or None,
        b["trips"], b["routes_n"], b["routes"], b["veh"], price,
        b["vexere_hotline"] or SUPPORT["vexere_hotline"], b["own_phones"], b["address"],
        b["url"], (b["note"] or b["status"] or "").strip(),
    ])
tot = ["", f"TOTAL — {len(BRANDS)} brands", None, f"=SUM(E5:E{4+len(rows)})", f"=SUM(F5:F{4+len(rows)})", None, "", "", "", "", "", "", "", ""]
nxt = write_table(ws, hdr, rows, number_cols={0, 2, 3, 4, 5}, totals=tot)

# hyperlinks for URL column
for i in range(len(rows)):
    r = 5 + i
    u = ws.cell(row=r, column=14).value
    if u:
        c = ws.cell(row=r, column=14)
        c.hyperlink = u
        c.font = Font(name=FONT_NAME, size=9, color=PRIMARY, underline="single")
        c.alignment = Alignment(horizontal="left", vertical="top")

nxt += 1
nxt = caption(ws, nxt, 2, "Rating/# Ratings = values shown on the brand's vexere.com page (JSON-LD aggregateRating). 'Vexere booking hotline' is the platform hotline shown for the brand; 'Operator's own phones' come from the page's contact blocks (văn phòng/điểm đón).", span=12)
nxt = caption(ws, nxt, 2, "10 brands without a dedicated vexere page (e.g. Cát Bà Limousine) or with a shared page are flagged in Notes; Quốc Phương Car's page is unreachable (redirect loop).", span=12)

auto_fit_columns(ws, min_width=6, max_width=24, header_row=4, data_start_row=5)
ws.column_dimensions["C"].width = 30
ws.column_dimensions["D"].width = 9
ws.column_dimensions["E"].width = 11
ws.column_dimensions["H"].width = 34
ws.column_dimensions["I"].width = 22
ws.column_dimensions["J"].width = 24
ws.column_dimensions["L"].width = 34
ws.column_dimensions["M"].width = 30
ws.column_dimensions["N"].width = 34
ws.column_dimensions["O"].width = 24
ws.freeze_panes = "D5"

# ============================================================ 4. SUPPORT PHONES
ws = wb.create_sheet("Support Phones")
setup_sheet(ws, title="Operator Support Phone Numbers — Web-Researched (Top 30 Brands)", last_col=8)
ws.cell(row=3, column=2, value="Official hotlines found via web search (official websites, fanpages, booking sites) — cross-checked against vexere.com brand pages").font = font_caption()
ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=8)

top = [b for b in BRANDS if b["web_phones"]]
rows = []
for i, b in enumerate(top, 1):
    rows.append([
        i, b["name"], b["trips"], b["web_phones"], b["web_website"],
        b["own_phones"] or "—", b["vexere_hotline"] or SUPPORT["vexere_hotline"], b["web_source"],
    ])
hdr = ["#", "Brand", "Trips (dataset)", "Operator support hotlines (web research)", "Official website",
       "Phones listed on vexere page", "Vexere booking hotline", "Source"]
nxt = write_table(ws, hdr, rows, number_cols={0, 2})

nxt += 1
nxt = caption(ws, nxt, 2, f"All other 276 brands can be booked via Vexere's nationwide hotline {SUPPORT['vexere_hotline']} ({SUPPORT['vexere_hours']}) — shown per brand on its vexere page; 95 brands additionally list their own office phones there (see Brand Directory).", span=7)
nxt = caption(ws, nxt, 2, "Hotline fee note: 1900-prefix numbers are charged per minute (typically 1,000đ); operator direct lines (03x/07x/08x/09x) are standard mobile rates. *1516 = short-code call from in-network SIM.", span=7)

auto_fit_columns(ws, min_width=6, max_width=24, header_row=4, data_start_row=5)
ws.column_dimensions["C"].width = 26
ws.column_dimensions["E"].width = 48
ws.column_dimensions["F"].width = 20
ws.column_dimensions["G"].width = 36
ws.column_dimensions["I"].width = 32
ws.freeze_panes = "C5"

wb.save(OUT)
print("PART 1 SAVED:", OUT)
print("sheets so far:", wb.sheetnames)
