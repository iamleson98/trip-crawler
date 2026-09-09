#!/usr/bin/env python3
"""Fix Review sheet cross-check formulas: tighten ranges to actual data extents
(exclude totals rows), correct distribution column B→C."""
from openpyxl import load_workbook

F = "/home/z/my-project/download/vexere_brands_feedback_report.xlsx"
wb = load_workbook(F)
ws = wb["Review"]

FIX = {
    5: "=COUNTA('Brand Directory'!C5:C310)",        # 306 data rows
    6: "=COUNTA('Customer Feedback'!C5:C1203)",     # 1199 review rows
    7: "=COUNTA('Support Phones'!C5:C34)",          # 30 researched brands
    8: "=COUNTA('Vexere Features'!C5:C28)",         # 24 features
    9: "=SUM('Brand Directory'!E5:E310)",           # cumulative ratings, excl. totals row
    10: "=SUM('Ratings Analysis'!C5:C11)",          # distribution bucket counts (col C)
}
for r, f in FIX.items():
    ws.cell(row=r, column=4, value=f)

wb.save(F)
print("Review formulas fixed:", FIX)
